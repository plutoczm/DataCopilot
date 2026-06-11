from pathlib import Path
from uuid import uuid4

from docx import Document

from backend.app.application.rag.chunking_service import clean_text
from backend.app.application.rag.models import LoadedDocument


class DOCXLoader:
    file_type = "docx"

    def load(
        self,
        path: Path,
        *,
        domain: str,
        tags: list[str] | None = None,
    ) -> LoadedDocument:
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return LoadedDocument(
            document_id=uuid4().hex,
            filename=path.name,
            file_type=self.file_type,
            domain=domain,
            source_path=path,
            stored_path=path,
            content=clean_text(text),
            tags=tags or [],
        )

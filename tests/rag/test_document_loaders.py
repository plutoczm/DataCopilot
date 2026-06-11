from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfWriter

from backend.app.infrastructure.document_loaders import DocumentLoaderFactory


def test_txt_loader_reads_and_cleans_text(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Spark   AQE\n\n\nreduces shuffle.", encoding="utf-8")

    loaded = DocumentLoaderFactory().load(path, domain="spark")

    assert loaded.filename == "notes.txt"
    assert loaded.file_type == "txt"
    assert loaded.domain == "spark"
    assert loaded.content == "Spark AQE\n\nreduces shuffle."


def test_markdown_loader_preserves_meaningful_text(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text("# Kafka\n\n- Consumer lag\n- Rebalance", encoding="utf-8")

    loaded = DocumentLoaderFactory().load(path, domain="kafka")

    assert loaded.file_type == "markdown"
    assert "Kafka" in loaded.content
    assert "Consumer lag" in loaded.content


def test_docx_loader_reads_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "warehouse.docx"
    document = DocxDocument()
    document.add_paragraph("ODS stores raw business data.")
    document.add_paragraph("DWD cleans and standardizes it.")
    document.save(path)

    loaded = DocumentLoaderFactory().load(path, domain="warehouse")

    assert loaded.file_type == "docx"
    assert "ODS stores raw business data." in loaded.content
    assert "DWD cleans and standardizes it." in loaded.content


def test_pdf_loader_handles_empty_pdf_without_crashing(tmp_path: Path) -> None:
    path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as file:
        writer.write(file)

    loaded = DocumentLoaderFactory().load(path, domain="linux")

    assert loaded.file_type == "pdf"
    assert loaded.filename == "empty.pdf"
    assert loaded.content == ""


def test_loader_factory_rejects_unsupported_file_type(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2", encoding="utf-8")

    try:
        DocumentLoaderFactory().load(path, domain="misc")
    except ValueError as exc:
        assert "Unsupported file type" in str(exc)
    else:
        raise AssertionError("Unsupported file type was accepted")
